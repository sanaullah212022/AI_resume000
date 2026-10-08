"""ATS Resume Checker - Streamlit app powered by Google Gemini Flash.

Upload a resume (PDF, DOCX or TXT), optionally paste a job description, and get
an ATS score, a section-by-section breakdown and concrete improvements.
"""

from __future__ import annotations

import io
import json
import os
import re
import time

import streamlit as st

# --------------------------------------------------------------------------- #
# Config
# --------------------------------------------------------------------------- #
DEFAULT_MODEL = "gemini-3.5-flash"
MODEL_CHOICES = ["gemini-3.5-flash", "gemini-3-flash-preview", "gemini-2.5-flash"]
MAX_RESUME_CHARS = 20_000
MAX_JD_CHARS = 8_000
MIN_RESUME_CHARS = 150

st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")


# --------------------------------------------------------------------------- #
# Text extraction
# --------------------------------------------------------------------------- #
def extract_text_from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            unlocked = reader.decrypt("")  # returns 0 when the password is wrong
        except Exception:
            unlocked = 0
        if not unlocked:
            raise ValueError("This PDF is password protected. Please upload an unlocked copy.")
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n".join(pages)


def extract_text_from_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_resume_text(uploaded_file) -> str:
    """Return plain text from an uploaded PDF / DOCX / TXT file."""
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()
    if name.endswith(".pdf"):
        text = extract_text_from_pdf(data)
    elif name.endswith(".docx"):
        text = extract_text_from_docx(data)
    elif name.endswith(".txt"):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file type. Upload a PDF, DOCX or TXT file.")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return text


# --------------------------------------------------------------------------- #
# Quick local (non-AI) checks - fast, free, deterministic
# --------------------------------------------------------------------------- #
SECTION_PATTERNS = {
    "Summary / Objective": r"\b(summary|objective|profile|about me)\b",
    "Experience": r"\b(experience|employment|work history|internship)s?\b",
    "Education": r"\b(education|academic|qualification)s?\b",
    "Skills": r"\b(skills|technologies|technical skills|competencies)\b",
    "Projects": r"\b(projects?|portfolio)\b",
    "Certifications": r"\b(certifications?|licenses?|courses?)\b",
}


def local_checks(text: str) -> dict:
    lower = text.lower()
    words = re.findall(r"\b\w+\b", text)
    return {
        "word_count": len(words),
        "email": bool(re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)),
        "phone": bool(re.search(r"(\+?\d[\d\s().-]{8,}\d)", text)),
        "linkedin": "linkedin.com" in lower,
        "github": "github.com" in lower,
        "sections": {
            name: bool(re.search(pattern, lower)) for name, pattern in SECTION_PATTERNS.items()
        },
        "has_numbers": len(re.findall(r"\b\d+(?:\.\d+)?%?", text)) >= 5,
    }


# --------------------------------------------------------------------------- #
# Gemini
# --------------------------------------------------------------------------- #
PROMPT_TEMPLATE = """You are an expert technical recruiter and Applicant Tracking System (ATS) analyst.
Evaluate the resume below the way a strict ATS parser plus a human recruiter would.

{jd_block}

Scoring guidance (be honest and realistic, do not inflate scores):
- keywords: relevant hard skills, tools and role keywords{jd_hint}
- formatting: clean structure, standard section headings, parseable text, consistent dates
- experience: impact-focused bullets, action verbs, quantified results
- education_skills: clear education and a well organised skills section
- readability: concise, no typos, appropriate length

Return ONLY valid JSON (no markdown, no commentary) with exactly this structure:
{{
  "ats_score": <integer 0-100>,
  "summary": "<2-3 sentence overall assessment>",
  "section_scores": {{
    "keywords": <integer 0-100>,
    "formatting": <integer 0-100>,
    "experience": <integer 0-100>,
    "education_skills": <integer 0-100>,
    "readability": <integer 0-100>
  }},
  "strengths": ["<short point>", "..."],
  "weaknesses": ["<short point>", "..."],
  "missing_keywords": ["<keyword>", "..."],
  "found_keywords": ["<keyword>", "..."],
  "improvements": [
    {{"priority": "High|Medium|Low", "area": "<section or topic>", "suggestion": "<specific, actionable fix>"}}
  ],
  "bullet_rewrites": [
    {{"original": "<weak bullet copied from the resume>", "improved": "<stronger rewritten bullet>"}}
  ]
}}

Rules:
- Give 5-10 improvements ordered by priority, and 3-5 bullet rewrites taken from the resume.
- Never invent experience, employers, degrees or metrics. In rewrites, use placeholders like [X%] where a number is needed.
- missing_keywords should be keywords a recruiter would expect for this profile{jd_missing_hint}.

RESUME:
\"\"\"
{resume}
\"\"\"
"""


def build_prompt(resume_text: str, job_description: str) -> str:
    jd = job_description.strip()
    if jd:
        jd_block = f'JOB DESCRIPTION (score the resume against this role):\n"""\n{jd[:MAX_JD_CHARS]}\n"""'
        jd_hint = " and how well they match the job description"
        jd_missing_hint = " that appear in the job description but not in the resume"
    else:
        jd_block = "No job description was provided. Judge the resume for its apparent target role in general."
        jd_hint = ""
        jd_missing_hint = ""
    return PROMPT_TEMPLATE.format(
        jd_block=jd_block,
        jd_hint=jd_hint,
        jd_missing_hint=jd_missing_hint,
        resume=resume_text[:MAX_RESUME_CHARS],
    )


def parse_model_json(raw: str) -> dict:
    """Parse JSON from the model, tolerating code fences or stray text."""
    if not raw:
        raise ValueError("The model returned an empty response.")
    cleaned = raw.strip()
    cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned)
    cleaned = re.sub(r"\s*```$", "", cleaned)
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        start, end = cleaned.find("{"), cleaned.rfind("}")
        if start != -1 and end > start:
            return json.loads(cleaned[start : end + 1])
        raise ValueError("Could not read the model's response as JSON.")


def _clamp(value, default=0) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return default


def _str_list(value) -> list[str]:
    if not isinstance(value, list):
        return []
    return [str(v).strip() for v in value if str(v).strip()]


def normalize_result(data: dict) -> dict:
    """Make sure every field exists and has the right type, so the UI never crashes."""
    section_scores = data.get("section_scores") if isinstance(data.get("section_scores"), dict) else {}
    improvements = []
    for item in data.get("improvements") or []:
        if isinstance(item, dict):
            priority = str(item.get("priority", "Medium")).strip().capitalize()
            if priority not in ("High", "Medium", "Low"):
                priority = "Medium"
            improvements.append(
                {
                    "priority": priority,
                    "area": str(item.get("area", "General")).strip(),
                    "suggestion": str(item.get("suggestion", "")).strip(),
                }
            )
        elif isinstance(item, str) and item.strip():
            improvements.append({"priority": "Medium", "area": "General", "suggestion": item.strip()})
    rewrites = []
    for item in data.get("bullet_rewrites") or []:
        if isinstance(item, dict) and item.get("improved"):
            rewrites.append(
                {"original": str(item.get("original", "")).strip(), "improved": str(item["improved"]).strip()}
            )
    return {
        "ats_score": _clamp(data.get("ats_score")),
        "summary": str(data.get("summary", "")).strip(),
        "section_scores": {
            "Keywords": _clamp(section_scores.get("keywords")),
            "Formatting": _clamp(section_scores.get("formatting")),
            "Experience": _clamp(section_scores.get("experience")),
            "Education & Skills": _clamp(section_scores.get("education_skills")),
            "Readability": _clamp(section_scores.get("readability")),
        },
        "strengths": _str_list(data.get("strengths")),
        "weaknesses": _str_list(data.get("weaknesses")),
        "missing_keywords": _str_list(data.get("missing_keywords")),
        "found_keywords": _str_list(data.get("found_keywords")),
        "improvements": improvements,
        "bullet_rewrites": rewrites,
    }


def analyze_resume(api_key: str, model: str, resume_text: str, job_description: str) -> dict:
    from google import genai
    from google.genai import types

    client = genai.Client(api_key=api_key)
    prompt = build_prompt(resume_text, job_description)
    config = types.GenerateContentConfig(
        temperature=0.2,
        response_mime_type="application/json",
    )

    last_error: Exception | None = None
    for attempt in range(2):  # one retry for transient errors / bad JSON
        try:
            response = client.models.generate_content(model=model, contents=prompt, config=config)
            return normalize_result(parse_model_json(response.text))
        except Exception as exc:  # noqa: BLE001 - surfaced to the user below
            last_error = exc
            message = str(exc).lower()
            # Do not retry on auth / model-name problems
            if any(k in message for k in ("api key", "api_key", "permission", "not found", "invalid argument", "401", "403", "404")):
                break
            time.sleep(1.5)
    raise RuntimeError(friendly_error(last_error))


def friendly_error(exc: Exception | None) -> str:
    text = str(exc) if exc else "Unknown error"
    low = text.lower()
    if "api key" in low or "api_key" in low or "401" in low or "403" in low or "permission" in low:
        return "Gemini rejected the API key. Check that it is correct and has access to the Gemini API."
    if "not found" in low or "404" in low:
        return "That model name was not found. Pick another model in the sidebar."
    if "429" in low or "quota" in low or "rate limit" in low or "resource_exhausted" in low:
        return "Rate limit or quota reached. Wait a minute and try again."
    return f"Analysis failed: {text[:300]}"


# --------------------------------------------------------------------------- #
# UI helpers
# --------------------------------------------------------------------------- #
def score_label(score: int) -> tuple[str, str]:
    if score >= 80:
        return "Excellent", "green"
    if score >= 65:
        return "Good", "blue"
    if score >= 50:
        return "Needs work", "orange"
    return "Poor", "red"


def get_api_key(sidebar_value: str) -> str:
    if sidebar_value.strip():
        return sidebar_value.strip()
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return str(st.secrets["GEMINI_API_KEY"]).strip()
    except Exception:
        pass
    return os.environ.get("GEMINI_API_KEY", "").strip()


def render_local_checks(checks: dict) -> None:
    def mark(ok: bool) -> str:
        return "✅" if ok else "❌"

    c1, c2 = st.columns(2)
    with c1:
        st.markdown("**Contact & links**")
        st.write(f"{mark(checks['email'])} Email address")
        st.write(f"{mark(checks['phone'])} Phone number")
        st.write(f"{mark(checks['linkedin'])} LinkedIn URL")
        st.write(f"{mark(checks['github'])} GitHub URL (optional)")
        st.write(f"{mark(checks['has_numbers'])} Quantified achievements (numbers / %)")
        wc = checks["word_count"]
        length_ok = 300 <= wc <= 900
        st.write(f"{mark(length_ok)} Length: {wc} words (ideal 300-900)")
    with c2:
        st.markdown("**Sections detected**")
        for name, found in checks["sections"].items():
            st.write(f"{mark(found)} {name}")


def render_results(result: dict, checks: dict) -> None:
    score = result["ats_score"]
    label, color = score_label(score)

    top1, top2 = st.columns([1, 2])
    with top1:
        st.metric("ATS Score", f"{score} / 100")
        st.markdown(f"**Rating:** :{color}[{label}]")
        st.progress(score / 100)
    with top2:
        st.markdown("**Overall assessment**")
        st.write(result["summary"] or "No summary returned.")

    st.divider()
    st.subheader("Score breakdown")
    cols = st.columns(len(result["section_scores"]))
    for col, (name, value) in zip(cols, result["section_scores"].items()):
        col.metric(name, value)
        col.progress(value / 100)

    tab_fix, tab_kw, tab_rewrite, tab_sw, tab_checks = st.tabs(
        ["🛠 Improvements", "🔑 Keywords", "✍️ Bullet rewrites", "⚖️ Strengths & weaknesses", "📋 Quick checks"]
    )

    with tab_fix:
        if not result["improvements"]:
            st.info("No improvements returned.")
        order = {"High": 0, "Medium": 1, "Low": 2}
        icons = {"High": "🔴", "Medium": "🟠", "Low": "🟢"}
        for item in sorted(result["improvements"], key=lambda i: order[i["priority"]]):
            st.markdown(f"{icons[item['priority']]} **{item['priority']} · {item['area']}**")
            st.write(item["suggestion"])

    with tab_kw:
        k1, k2 = st.columns(2)
        with k1:
            st.markdown("**Missing keywords** (consider adding if truthful)")
            if result["missing_keywords"]:
                st.write(", ".join(f"`{k}`" for k in result["missing_keywords"]))
            else:
                st.write("None flagged.")
        with k2:
            st.markdown("**Keywords found**")
            if result["found_keywords"]:
                st.write(", ".join(f"`{k}`" for k in result["found_keywords"]))
            else:
                st.write("None returned.")

    with tab_rewrite:
        if not result["bullet_rewrites"]:
            st.info("No bullet rewrites returned.")
        for i, item in enumerate(result["bullet_rewrites"], 1):
            st.markdown(f"**Example {i}**")
            if item["original"]:
                st.markdown(f"Before: ~~{item['original']}~~")
            st.success(item["improved"])
        if result["bullet_rewrites"]:
            st.caption("Replace placeholders like [X%] with your real numbers. Never add claims that are not true.")

    with tab_sw:
        s1, s2 = st.columns(2)
        with s1:
            st.markdown("**Strengths**")
            for s in result["strengths"] or ["None returned."]:
                st.write(f"✅ {s}")
        with s2:
            st.markdown("**Weaknesses**")
            for w in result["weaknesses"] or ["None returned."]:
                st.write(f"⚠️ {w}")

    with tab_checks:
        render_local_checks(checks)

    report = {"ats_score": score, **{k: v for k, v in result.items() if k != "ats_score"}}
    st.download_button(
        "⬇️ Download report (JSON)",
        data=json.dumps(report, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


# --------------------------------------------------------------------------- #
# Main
# --------------------------------------------------------------------------- #
def main() -> None:
    st.title("📄 ATS Resume Checker")
    st.caption("Upload your resume to get an ATS score and specific ways to improve it. Powered by Google Gemini Flash.")

    with st.sidebar:
        st.header("Settings")
        api_key_input = st.text_input(
            "Gemini API key",
            type="password",
            help="Get a free key at https://aistudio.google.com/apikey. "
            "On Streamlit Cloud you can store it in Secrets as GEMINI_API_KEY instead.",
        )
        model = st.selectbox("Model", MODEL_CHOICES, index=MODEL_CHOICES.index(DEFAULT_MODEL))
        st.markdown("---")
        st.caption(
            "Your resume is sent to the Gemini API for analysis and is not stored by this app. "
            "Avoid uploading documents you are not comfortable sharing."
        )

    left, right = st.columns([1, 1])
    with left:
        uploaded = st.file_uploader("Upload resume", type=["pdf", "docx", "txt"])
    with right:
        job_description = st.text_area(
            "Job description (optional, recommended)",
            height=150,
            placeholder="Paste the job description here to get a role-specific score and keyword gap...",
        )

    analyze_clicked = st.button("Analyze resume", type="primary", disabled=uploaded is None)

    if analyze_clicked and uploaded is not None:
        api_key = get_api_key(api_key_input)
        if not api_key:
            st.error("Please enter your Gemini API key in the sidebar.")
            st.stop()

        try:
            with st.spinner("Reading your resume..."):
                resume_text = extract_resume_text(uploaded)
        except Exception as exc:  # noqa: BLE001
            st.error(f"Could not read the file: {exc}")
            st.stop()

        if len(resume_text) < MIN_RESUME_CHARS:
            st.error(
                "Very little text could be extracted. If your PDF is a scanned image, "
                "an ATS cannot read it either - export a text-based PDF or upload a DOCX."
            )
            st.stop()

        try:
            with st.spinner("Analyzing with Gemini..."):
                result = analyze_resume(api_key, model, resume_text, job_description)
        except Exception as exc:  # noqa: BLE001
            st.error(str(exc))
            st.stop()

        st.session_state["result"] = result
        st.session_state["checks"] = local_checks(resume_text)

    if "result" in st.session_state:
        render_results(st.session_state["result"], st.session_state["checks"])
    elif uploaded is None:
        st.info("Upload a resume (PDF, DOCX or TXT) to get started.")


if __name__ == "__main__":
    main()
